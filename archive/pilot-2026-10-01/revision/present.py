"""Render the revised experiment in the original deck's plain style."""
from html import escape
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pilot.common import ROOT,read
from revision.messaging import audit_weather_slides
s=read(ROOT/'results/revision-summary.json');runs=s['runs']
findings=read(ROOT/'revision/findings.json') if (ROOT/'revision/findings.json').exists() else {}
ARM_COLORS={'scratch':'777777','rhiza':'2F6B8A','accord':'8A1C1C'}
ARM_LIGHT={'scratch':'CCCCCC','rhiza':'BCD2E0','accord':'DEB8B8'}
def text(sl,x,y,w,h,value,size=18,bold=False,color='111111'):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.margin_left=tf.margin_right=0;tf.word_wrap=True
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(10)
 return sh

def new(title,subtitle):
 sl=prs.slides.add_slide(prs.slide_layouts[0]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor(255,255,255)
 text(sl,.6,.45,12.1,.85,title,30,True);text(sl,.6,1.5,12.1,.8,subtitle)
 text(sl,.6,7,12.1,.3,f'ACCORD · 1 October 2026 · {len(prs.slides)}',10,color='6B6B6B')
 return sl

def table(sl,rows,y,widths,height=.53,size=16):
 t=sl.shapes.add_table(len(rows),len(rows[0]),Inches(.6),Inches(y),Inches(12.1),Inches(height*len(rows))).table
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for i,row in enumerate(rows):
  for j,v in enumerate(row):
   c=t.cell(i,j);c.text=str(v);c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string('EEEEEE' if i==0 else 'FFFFFF');c.margin_left=Inches(.09);c.margin_top=Inches(.05 if height<.5 else .08);c.margin_bottom=Inches(.04)
   for p in c.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(size);p.font.bold=i==0 or j==0;p.font.color.rgb=RGBColor.from_string(ARM_COLORS.get(str(v).split(' / ')[-1].lower(),'111111') if j==0 and i>0 else '111111')

def fmt(key,field):
 if field=='rps':
  v=runs.get(key,{}).get('metrics',{}).get('RPS',{}).get('submitted')
  return f'{v:.3f}' if v is not None else '—'
 v=runs.get(key,{}).get(field)
 if v is None:return '—'
 if field=='usd':return f'${v:.2f}'+('+' if runs[key]['usage_incomplete'] else '')
 if field=='seconds':return f'{v/60:.1f} min'
 if field=='tokens':return f'{v:,}'
 return str(v)

def rect(sl,x,y,w,h,fill='EEEEEE'):
 sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background()
 return sh

initials=[('haiku-scratch-initial','Haiku / scratch'),('haiku-rhiza-initial','Haiku / Rhiza'),('haiku-accord-initial','Haiku / ACCORD'),('sonnet-scratch-initial','Sonnet / scratch'),('sonnet-rhiza-initial','Sonnet / Rhiza'),('sonnet-accord-initial','Sonnet / ACCORD'),('fable-scratch-initial','Fable / scratch'),('fable-rhiza-initial','Fable / Rhiza'),('fable-accord-initial','Fable / ACCORD')]

def build(use_tables=False):
 global prs
 prs=Presentation(ROOT/'slides/checkin-original.pptx')
 for slide in list(prs.slides)[9:15]:
  for sh in slide.shapes:
   if sh.has_text_frame:
    for p in sh.text_frame.paragraphs:
     for r in p.runs:r.text=r.text.replace('Current experiment:','Earlier experiment:').replace('Current task:','Earlier task:').replace('Next: test the libraries','Earlier proposal: test the libraries')
 audit_weather_slides(prs,text,ARM_COLORS)
 sl=new('The task: a reusable seasonal rainfall forecast',
        'Same data and budgets; scratch, Rhiza skills, or ACCORD libraries (DeepScale / Rosetta).')
 text(sl,.6,2.15,6.1,.95,'Turn September-issued ECMWF forecasts and CHIRPS observations into Oct–Dec rainfall probabilities for a 12-cell Kenyan study box.',20)
 text(sl,.6,3.4,6.1,.9,'Train on 1993–2004; predict 2005–06. Save the code and fitted model, then forecast 2007–08 in a fresh conversation.',20)
 text(sl,.6,4.65,6.1,.9,'Check completion and replay, scientific quality with an LLM judge, forecast error, tokens, time and cost.',20)
 text(sl,.6,5.95,6.1,.85,'Can stronger agents finish? Do tools save effort? Does saved work help next time?',19,True)
 sl.shapes.add_picture(str(ROOT/'slides/assets/tercile-example.png'),Inches(6.8),Inches(2.15),width=Inches(6.05))
 text(sl,7.1,6.65,5.5,.24,'Example: Fable + ACCORD, historical 2005 forecast; 33% is climatology.',10,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='The Kenya locator places the shaded study box beside its probability grid. Natural Earth generalized boundary, source/hash in slides/assets/kenya-study-area.json. Forecast map source and hash: slides/assets/tercile-example.json. Above-normal rainfall is above the observed 1993–2004 upper-tercile threshold; this is one of three probabilities, not a map of skill. The 12 cells cover south=-3, north=1, west=36, east=39, not a national forecast. Monthly rates are converted to seasonal totals; CHIRPS observations are spatially aggregated with cosine-latitude weights. Leave-one-year-out validation fits both thresholds and calibration within each fold. Forecast-only frozen inference receives no verification observations. Twelve independent artifact, calculation and replay checks plus five LLM-judge criteria; accuracy uses RPS/RMSE separately from completion and cost. These are previously inspected development years, two per batch, so general skill and calibration are not established. Same per-attempt caps: $5.50, 900 seconds, 800k tokens, 60 turns, 16,384 output tokens. Data are supplied; acquisition is not tested.'

 sl=new('Model capability and tools', 'Nine initial attempts: three models × three tool environments, with the same task, data and limits.')
 for x,arm,label in [(0.6,'scratch','Scratch'),(2.4,'rhiza','Rhiza'),(4.0,'accord','ACCORD')]:
  rect(sl,x,2.05,.18,.18,ARM_COLORS[arm]);text(sl,x+.27,1.97,1.4,.3,label,14)
 if use_tables:
  rows=[['Model / tools','Checks','Tokens','Minutes','API cost','RPS ↓']]
  for key,label in initials:
   rows.append([label,fmt(key,'checks'),fmt(key,'tokens'),f"{runs[key]['seconds']/60:.2f}",fmt(key,'usd'),fmt(key,'rps')])
  table(sl,rows,2.35,[3.0,1.2,2.2,1.5,1.6,2.6],height=.35,size=14)
 else:
  for x,w,label in [(.6,2.6,'Model / tools'),(3.05,2.4,'Checks / 12'),(6.1,2.6,'Agent API cost ($)'),(9.72,1.3,'Tokens'),(11.3,1.2,'Minutes')]:text(sl,x,2.37,w,.35,label,15,True)
  for i,(key,label) in enumerate(initials):
   y=2.87+i*.33+(i//3)*.10;r=runs[key];n=int(str(r['checks']).split('/')[0]);accent=ARM_COLORS[r['arm']]
   text(sl,.6,y-.03,2.45,.31,label,15)
   rect(sl,3.05,y+.03,2.0,.20,'EEEEEE');rect(sl,3.05,y+.03,2.0*n/12,.20,accent)
   text(sl,5.2,y-.03,.75,.31,str(n)+'/12',14)
   rect(sl,6.1,y+.03,2.8,.20,'EEEEEE');rect(sl,6.1,y+.03,2.8*r['usd']/5,.20,accent)
   text(sl,9.0,y-.03,.7,.31,fmt(key,'usd'),14)
   text(sl,9.72,y-.03,1.4,.31,fmt(key,'tokens'),15)
   text(sl,11.3,y-.03,1.1,.31,f"{r['seconds']/60:.1f}",15)
 text(sl,.6,6.05,12.1,.5,findings.get('matrix','Full matrix evaluation pending.'),17)
 text(sl,.6,6.65,12.1,.25,'One attempt per condition. Cost bars: $0–$5. Haiku 4.5 / Sonnet 5.5 / Fable 5.1. Rhiza supplied catalog context; scripts were not called.',11,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='Colors are arm-specific on all comparison charts: scratch gray, Rhiza blue, ACCORD red. Rows are grouped Haiku / Sonnet / Fable, always scratch / Rhiza / ACCORD within each group. Haiku + ACCORD and Sonnet + ACCORD are separately authorized extensions with identical task, model settings, pinned sandbox image, tool guidance and inputs. No retries or alternate prompts. All attempts retained. Sonnet scratch missed three units-metadata checks despite correct array values. Fable Rhiza has a read-only replay defect. Rhiza conditions did not call scripts or read individual SKILL.md files. ACCORD agents use the installed scientific API. Checks and exact RPS/RMSE, model versions, prompts, traces, tokens and USD are in report.html and results/revision-summary.json. Sonnet + Rhiza matched Fable scratch’s 12/12 checks with 49% fewer tokens, 67% less elapsed time and 89% lower agent API cost. This does not mean skills improve every model or that libraries universally improve accuracy.'

 sl=new('Follow-up task: forecast the next two years',
        'Produce Oct–Dec rainfall probabilities for 2007–08 in the same Kenyan study area (initial: 2005–06).\nFable, fresh conversation: reuse the saved model without refitting; validate the workflow and record repairs.')
 if use_tables:
  rows=[['Fable arm','Initial → follow-up $','Follow-up tokens','Minutes','Checks']]
  for arm,label in [('scratch','Scratch'),('rhiza','Rhiza'),('accord','ACCORD')]:
   k=f'fable-{arm}-followup';rows.append([label,fmt(f'fable-{arm}-initial','usd')+' → '+fmt(k,'usd'),fmt(k,'tokens'),f"{runs[k]['seconds']/60:.2f}",fmt(k,'checks')])
  table(sl,rows,2.65,[2.2,3.2,2.6,2.3,1.8])
 else:
  text(sl,.6,2.55,3.3,.32,'Agent API cost ($0–$5)',15,True)
  text(sl,3.9,2.55,3.05,.32,'Light: initial · dark: follow-up',12)
  text(sl,8.05,2.55,4.65,.35,'Follow-up: tokens / time / checks',15,True)
  for i,(arm,label) in enumerate([('scratch','Scratch'),('rhiza','Rhiza'),('accord','ACCORD')]):
   y=3.15+i*.82;text(sl,.6,y+.1,1.5,.35,label,18,True,color=ARM_COLORS[arm])
   for offset,phase,color_ in [(0,'initial',ARM_LIGHT[arm]),(.29,'followup',ARM_COLORS[arm])]:
    key=f'fable-{arm}-{phase}';v=runs[key]['usd'];rect(sl,2.25,y+offset,4.7*v/5,.2,color_);text(sl,2.36+4.7*v/5,y+offset-.08,.85,.34,fmt(key,'usd'),14)
   k=f'fable-{arm}-followup';text(sl,8.05,y+.16,4.65,.4,fmt(k,'tokens')+' / '+fmt(k,'seconds')+' / '+fmt(k,'checks'),17)
 text(sl,.6,5.9,12.1,.5,'ACCORD’s follow-up: 50% fewer tokens, 50% lower cost and 31% less time than scratch.\nBoth passed 12/12 checks. Across both attempts, ACCORD cost 12% less.',17)
 text(sl,.6,6.65,12.1,.25,'RPS ↓: scratch 0.396; Rhiza 0.404; ACCORD 0.412. No accuracy advantage established. No fresh-workspace controls or download savings.',11,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='Extension task: with the same 12 cells and training data, prepare the saved predictor for the September-issued October–December forecasts for 2007 and 2008, instead of the initial 2005 and 2006. Each Fable arm starts a fresh conversation with its own retained files. No refitting; repairs may be documented. Verification observations remain private. The locator on slide 16 marks the same study box for both phases. Follow-up uses 2007–08, two previously inspected development years. All three Fable arms retained unchanged fitted-state hashes and cached inputs. Direct predictors execute in approximately 1–2 seconds without LLM calls. Bars instead measure conversational follow-up including validation and reporting. ACCORD initial + follow-up cost $6.27595 versus scratch $7.13088. No causal reuse estimate without fresh-workspace controls. Sonnet has initial comparisons only; follow-ups remain Fable only.'

 sl=new('How good were the forecasts?',
        'Fable in all three arms, evaluated against CHIRPS observations.\nRanked Probability Score (RPS): error in below / near / above probabilities; lower is better.')
 quality_rows=[['Forecast','2005–06 initial','2007–08 follow-up']]
 for label,baseline in [('Raw ECMWF','raw'),('Simple model-climatology baseline','model_climate')]:
  values=[]
  for phase in ['initial','followup']:
   value=runs[f'fable-scratch-{phase}']['metrics']['RPS'][baseline]
   assert all(abs(runs[f'fable-{arm}-{phase}']['metrics']['RPS'][baseline]-value)<1e-12 for arm in ['rhiza','accord'])
   values.append(f'{value:.3f}')
  quality_rows.append([label,*values])
 for arm,label in [('scratch','Fable / scratch'),('rhiza','Fable / Rhiza'),('accord','Fable / ACCORD')]:
  quality_rows.append([label,fmt(f'fable-{arm}-initial','rps'),fmt(f'fable-{arm}-followup','rps')])
 table(sl,quality_rows,2.5,[6.1,3.0,3.0],height=.44,size=17)
 text(sl,.6,5.45,12.1,.8,'All three beat raw ECMWF; initial forecast accuracy is similar across arms.\nThe simple baseline wins on the follow-up. No ACCORD accuracy advantage established.',18)
 text(sl,.6,6.6,12.1,.3,'Two previously inspected years per batch; general skill and calibration remain unproven. Rhiza’s predictor scored, but workflow replay failed.',11,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='Source: results/revision-summary.json, metrics.RPS for Fable scratch, Rhiza and ACCORD initial and follow-up attempts. All scores come from frozen predictors evaluated against private-to-agent CHIRPS observations on the same 12 cells. Training years 1993–2004; initial 2005–06; follow-up 2007–08. Raw baseline counts ECMWF members in observed-training terciles. The model-climatology baseline counts members against terciles from the model’s own historical ensemble distribution; it accounts for distribution differences without the agent-fitted regression. Lower RPS is better. Initial RPS: raw 0.368546598, model-climatology 0.295942372, scratch 0.279028737, Rhiza 0.281017925, ACCORD 0.279452220. Follow-up: raw 0.439242635, model-climatology 0.353362053, scratch 0.396182983, Rhiza 0.403827736, ACCORD 0.412328401. Rainfall-amount RMSE, initial/follow-up mm: raw 214.898/80.597, scratch 178.094/70.096, Rhiza 178.485/73.325, ACCORD 175.870/69.915. Those amount-error reductions do not establish an ACCORD advantage over scratch. Only two development years per batch; spatial cells are not independent years, and verification years had previously been inspected by the designer. Reliability diagnostics are descriptive, not evidence of generally calibrated probabilities. Rhiza’s frozen predictor produced scoreable outputs even though solve.py failed the separate full-workflow replay gate. No submissions were repaired, no new forecasts run, and no new model calls were made for this slide. The next scientific evaluation should apply saved predictors to more untouched years and compare against the same baselines.'

 sl=new('Toward auto-science', 'Make each experiment leave behind something the next agent can use.')
 text(sl,.6,2.65,12.1,.65,'Research  →  Hypothesis  →  Experiment  →  Evaluate',27,True)
 text(sl,.6,4.0,12.1,.7,'Keep tested code, data, fitted models, findings and failed attempts.',23)
 text(sl,.6,5.0,12.1,.7,'Use that shared record to choose and run the next experiment.',23)
 text(sl,.6,6.3,12.1,.4,'Today: measure completion, cost and reuse. Next: test whether scientific improvements accumulate.',16,color='555555')
 sl.notes_slide.notes_text_frame.text='Simplified from /Users/ezekielbarnett/Developer/ACCORD/docs/presentations/toward-auto-science.html and assets/slide-autoscience-loop.svg. The source describes retrieve/reason/experiment workers submitting to a shared evaluator, curating findings, dead ends, wiki pages and writeups, and using that record for the next hypothesis. This slide is the intended research direction; the weather pilot does not yet demonstrate recursive improvement, a shared scientific leaderboard, or community-contributed discoveries. No illustrative skill curve is presented as measured evidence.'
 if not use_tables:
  prs.save(ROOT/'2026-10-01-accord-checkin.pptx')
  weather=Presentation(ROOT/'2026-10-01-accord-checkin.pptx')
  for n,it in reversed(list(enumerate(list(weather.slides._sldIdLst),1))):
   if n not in [10,15,16,17,18,19,20]:weather.part.drop_rel(it.rId);weather.slides._sldIdLst.remove(it)
  weather.save(ROOT/'slides/weather-slides.pptx')
 for item in list(prs.slides._sldIdLst)[:15]:prs.part.drop_rel(item.rId);prs.slides._sldIdLst.remove(item)
 prs.save(ROOT/'slides'/('revised-tables.pptx' if use_tables else 'revised-slides.pptx'))

build()
build(use_tables=True)

rows=[];reviews=[]
for key,r in runs.items():
 metrics=r['metrics'];rps=metrics.get('RPS',{}).get('submitted');rmse=metrics.get('RMSE_mm',{}).get('submitted')
 vals=[key,r['checks'],f"{r['tokens']:,}",f"{r['seconds']/60:.2f}",f"${r['usd']:.3f}",f'{rps:.4f}' if rps is not None else '—',f'{rmse:.1f}' if rmse is not None else '—',str(r['audit']['truncated_responses'])]
 rows.append('<tr>'+''.join('<td>'+escape(v)+'</td>' for v in vals)+'</tr>')
 run=ROOT/'runs/revision-v1'/key;j=read(run/'judge.json') if (run/'judge.json').exists() else None
 ratings=''.join('<li><b>'+escape(k)+f" {v['score']}/2</b>: "+escape(v['reason'])+'</li>' for k,v in j['ratings'].items()) if j else '<li>Judge pending</li>'
 diagnostic=('<p><b>Independent diagnostic:</b> '+escape(read(run/'failure-diagnostics.json').get('interpretation','See diagnostic file for details.'))+f' <a href="runs/revision-v1/{key}/failure-diagnostics.json">Details</a></p>') if (run/'failure-diagnostics.json').exists() else ''
 reviews.append(f'<details><summary>{escape(key)} — {escape(r["completion"])}</summary><p>Failed checks: {escape(", ".join(r["failed_checks"]) or "none")}</p>{diagnostic}<p><a href="runs/revision-v1/{key}/frozen/">Code, fitted model and figures</a> · <a href="runs/revision-v1/{key}/checks.json">Accuracy/calibration metrics</a> · <a href="runs/revision-v1/{key}/audit.json">Toolkit/cache audit</a> · <a href="runs/revision-v1/{key}/run.json">Model and usage</a></p><ul>{ratings}</ul></details>')
page='''<!doctype html><meta charset="utf-8"><title>ACCORD revised comparison</title><style>body{font:16px/1.5 Calibri,Arial,sans-serif;max-width:1250px;margin:40px auto;padding:0 24px;color:#111}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:9px;border:1px solid #ccc;text-align:left}th{background:#eee}details{padding:14px 0;border-top:1px solid #ccc}summary{cursor:pointer;font-weight:bold}a{color:#245b8e}li{margin:10px 0}</style><h1>ACCORD: revised seasonal comparison</h1><p><a href="2026-10-01-accord-checkin.pptx">Updated check-in deck</a> · <a href="slides/revised-slides.pptx">Five-slide extract</a> · <a href="slides/revised-tables.pptx">Table alternative</a> · <a href="results/revision-summary.json">Recorded results</a> · <a href="revision/PLAN.txt">Frozen protocol</a></p>'''
page+=''.join('<p>'+escape(findings.get(k,'Evaluation in progress.'))+'</p>' for k in ['capability','tools','sonnet_accord','haiku_accord','reuse'])
page+='<table><tr><th>Attempt</th><th>Checks</th><th>Tokens</th><th>Minutes</th><th>Agent $</th><th>RPS ↓</th><th>RMSE mm ↓</th><th>Truncated responses</th></tr>'+''.join(rows)+'</table>'
page+=f'<p>New agent cost: ${s["new_agent_usd"]:.4f}; new judge cost: ${s["new_judge_usd"]:.4f}; prior recorded cost: ${s["prior_recorded_usd"]:.4f}. <b>Total recorded: ${s["total_recorded_usd"]:.4f}</b>. '+escape(s['cost_note'])+'</p>'
page+='<p>Same supplied monthly data and scientific task across conditions. Train on 1993–2004; frozen predictors receive 2005–06 and then 2007–08, never verification observations. These were previously inspected development years; RPS, RMSE and reliability diagnostics describe this small example, not general forecast skill. All initial and follow-up attempts have the same caps. This is not a repeated-seed model ranking. Rhiza and ACCORD have improved tool navigation; results must not be pooled with smoke-v3. All arms can retain code and cached data.</p>'
page+='<p>Tool use audit: the Rhiza conditions received the catalog descriptions but did not invoke skill scripts or open individual SKILL.md files. ACCORD did call the installed EnsembleRegressionMethod for fitting and prediction. No live retrieval was tested. The strict overall completion label remains partial if the judge lists any concern, even when all numerical checks and all five ratings pass. Judge statements are evidence to review, not ground truth; Sonnet scratch’s three gate failures were units metadata, not incorrect array values or negative rainfall.</p>'
if (ROOT/'results/revision-direct-reuse.json').exists():
 page+='<h2>Direct reuse (no LLM call)</h2>'
 for v in read(ROOT/'results/revision-direct-reuse.json'):
  page+=f'<p>{escape(v["run"])}: {v["seconds"]:.2f} seconds, zero agent tokens; '+escape(v['status'])+'.</p>'
page+='<h2>Checks and evidence-citing judge</h2>'+''.join(reviews)
(ROOT/'report.html').write_text(page)
print('Revised deck and report generated')
