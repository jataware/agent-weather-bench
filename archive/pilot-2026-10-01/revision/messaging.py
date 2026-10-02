"""Keep the original deck's earlier benchmark distinct from this measured pilot."""
from pptx.util import Inches

def audit_weather_slides(prs,text,colors):
 # Preserve original slide numbers while excluding an unrelated older benchmark
 # from the presentation sequence. The source deck itself remains untouched.
 for number in [11,12,13,14]:
  sl=prs.slides[number-1];sl._element.set('show','0')
  for sh in sl.shapes:
   if not sh.has_text_frame or not sh.text.strip():continue
   if sh.top < Inches(1):
    title={11:'Archived benchmark: execution conditions',12:'Archived task: two-week Kenya heat outlook',13:'Archived benchmark: interim snapshot',14:'Archived benchmark: failure modes'}[number]
    p=sh.text_frame.paragraphs[0]
    if p.runs:
     p.runs[0].text=title
     for run in list(p.runs)[1:]:run.text=''
    else:p.text=title
   elif sh.top >= Inches(7):
    p=sh.text_frame.paragraphs[0]
    if p.runs:
     p.runs[0].text=f'Archived Kenya heat benchmark · different task and protocol · {number}'
     for run in list(p.runs)[1:]:run.text=''
  sl.notes_slide.notes_text_frame.text='Historical content imported from the original check-in deck. Different task, conditions, model matrix and execution limits from the seasonal rainfall pilot. Hidden from the slideshow to avoid mixing claims; retained for reference. Its figures are not included in the current comparison.'
 def clear(sl):
  for sh in list(sl.shapes):sh._element.getparent().remove(sh._element)
 sl=prs.slides[9];clear(sl)
 text(sl,.6,.45,12.1,.85,'What are we testing?',30,True)
 for y,value in [(1.7,'Can a coding agent build a working forecast from scratch, and how does that depend on the model?'),(2.95,'Does adding Rhiza weather skills or ACCORD libraries improve completion, time or cost?'),(4.2,'When the next forecast is needed, does retained code and fitted state reduce the work?'),(5.6,'Report completion, forecast error, tokens, time and cost separately. Treat the expected advantages as hypotheses.')]:text(sl,.6,y,12.1,.9,value,23)
 text(sl,.6,7,12.1,.3,'ACCORD · 1 October 2026 · 10',10,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='This is the motivation for the actual three-arm seasonal pilot. No claim that Rhiza cannot produce reusable code, that ACCORD already improves accuracy, or that an extra ACCORD skills condition was tested.'
 sl=prs.slides[14];clear(sl)
 text(sl,.6,.45,12.1,.85,'The three arms we actually tested',30,True)
 for y,arm,label,value in [(1.7,'scratch','Scratch','Standard scientific Python; the agent builds the workflow.'),(2.9,'rhiza','Rhiza skills','The same Python environment + Rhiza weather skills catalog and scripts.'),(4.1,'accord','ACCORD','The same Python environment + africas2s (DeepScale) and Rosetta/acmadDL, with API documentation.')]:
  text(sl,.6,y,2.25,.5,label,23,True,color=colors[arm]);text(sl,3.1,y,9.6,.9,value,21)
 text(sl,.6,5.55,12.1,.6,'Identical supplied data and budgets. Every arm can retain its own code, inputs and fitted models.',19)
 text(sl,.6,6.4,12.1,.5,'Observed use: Rhiza catalog context, no script calls; ACCORD agents called africas2s. Data retrieval was not tested.',14,color='555555')
 text(sl,.6,7,12.1,.3,'ACCORD · 1 October 2026 · 15',10,color='6B6B6B')
 sl.notes_slide.notes_text_frame.text='Supersedes the unexecuted libraries-versus-libraries-plus-guides proposal in the original deck. There is no fourth arm or bespoke ACCORD skill-guide experiment. Generic Python is allowed everywhere. The configured Rhiza arm gets the actual catalog and scripts, but these attempts did not read individual SKILL.md files or invoke scripts. The ACCORD arm gets installed packages, package docs and a generic navigation pointer to the persistent ensemble-regression API. All tested data were supplied; Rosetta acquisition and download caching are available but their benefits were not measured here. Model IDs, configurations, artifacts and actual toolkit imports are logged.'
