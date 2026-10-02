"""Render the simple three-slide extract to HTML for local review and PDF export."""
from html import escape
import base64
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_FILL
from pilot.common import ROOT
import sys
stem=sys.argv[1] if len(sys.argv)>1 else 'three-smoke-slides'
p=Presentation(ROOT/'slides'/f'{stem}.pptx')
S=914400

def inch(v):return float(v)/S

def color(font):
 try:return '#'+str(font.color.rgb)
 except (AttributeError,TypeError):return '#111111'

def paragraphs(tf):
 out=[]
 for par in tf.paragraphs:
  if not par.text:continue
  runs=list(par.runs);font=par.font
  size=font.size.pt if font.size else 18
  parts=[]
  for r in runs:
   f=r.font
   parts.append(f'<span style="font-size:{f.size.pt if f.size else size}pt;font-weight:{700 if (f.bold or font.bold) else 400};color:{color(f) if f.color.type else color(font)}">{escape(r.text)}</span>')
  out.append('<p>'+(''.join(parts) if runs else escape(par.text))+'</p>')
 return ''.join(out)
slides=[]
for slide in p.slides:
 shapes=[]
 for sh in slide.shapes:
  pos=f'left:{inch(sh.left)}in;top:{inch(sh.top)}in;width:{inch(sh.width)}in;height:{inch(sh.height)}in'
  if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
   blob=base64.b64encode(sh.image.blob).decode()
   shapes.append(f'<img style="position:absolute;{pos}" src="data:{sh.image.content_type};base64,{blob}">')
  elif sh.has_table:
   rows=[]
   for i,row in enumerate(sh.table.rows):
    cells=[]
    for j,c in enumerate(row.cells):
     padding=f'{inch(c.margin_top)}in {inch(c.margin_right)}in {inch(c.margin_bottom)}in {inch(c.margin_left)}in'
     cells.append(f'<td style="width:{inch(sh.table.columns[j].width)}in;padding:{padding};background:{"#eee" if i==0 else "white"}">{paragraphs(c.text_frame)}</td>')
    rows.append(f'<tr style="height:{inch(row.height)}in">'+''.join(cells)+'</tr>')
   shapes.append(f'<table style="{pos}">'+''.join(rows)+'</table>')
  elif sh.has_text_frame:
   background=''
   if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.fill.type == MSO_FILL.SOLID:
    background=f';background:#{sh.fill.fore_color.rgb}'
   shapes.append(f'<div class="text" style="{pos}{background}">{paragraphs(sh.text_frame)}</div>')
 slides.append('<section>'+''.join(shapes)+'</section>')
html='''<!doctype html><meta charset="utf-8"><title>ACCORD measured smoke</title><style>
@page { size: 13.333333in 7.5in; margin: 0; }
* { box-sizing: border-box; } body { margin:0; background:#ddd; font-family:Calibri,Arial,sans-serif; }
section { width:13.333333in;height:7.5in;position:relative;background:white;page-break-after:always;overflow:hidden;margin:0 auto 20px; }
.text,table { position:absolute; } .text { padding-top:.05in; } p { margin:0 0 12pt;line-height:1.1; }
table { border-collapse:collapse;table-layout:fixed; } td { border:1px solid #bbb;padding:.1in;vertical-align:top; } td p { margin:0; }
@media print { body {background:white;} section {margin:0;} }
</style>'''+''.join(slides)
(ROOT/'slides'/f'{stem}.html').write_text(html)
print('HTML preview written')
