"""Generate the review deck. Uses existing feasibility evidence; never runs models."""
import html
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]
SCORES=json.loads((ROOT/'results/development-scores.json').read_text())
PREFLIGHT=json.loads((ROOT/'results/sandbox-preflight.json').read_text())
slides=[
 dict(tag='THE RESEARCH QUESTION',title='Give coding agents a substrate that lasts',
      subtitle='The agent does the work in every arm. We test what helps that work accumulate.',
      cards=[('Claude from scratch','Generic scientific Python and internet access. The agent discovers data and builds its own workflow.'),
             ('Claude + Rhiza','Pinned skills, scripts and documentation. The agent composes an existing procedure into a solution.'),
             ('Claude + ACCORD','Pinned DeepScale and Rosetta. The agent can build on data adapters, scientific methods and fitted models.')],
      takeaway='Hypothesis: maintained scientific capabilities reduce repeated work and improve results across iterations.',
      note='Expected ordering is a hypothesis. Rhiza includes reusable code; all three arms may retain their own work. One follow-up cannot establish recursive improvement or community-scale effects.'),
 dict(tag='A FAIR EXPERIMENT',title='Change the tools. Hold the opportunity constant.',
      subtitle='Same model, task, token/spend limits, compute resources, source permissions and data credentials.',
      cards=[('1  Initial attempt','An isolated workspace. Scratch cannot import the other toolkits. Rhiza receives its catalog; ACCORD receives its libraries and documentation.'),
             ('2  Freeze and evaluate','Save code, outputs and provenance. Check numbers independently, rerun offline, then use one evidence-citing LLM judge.'),
             ('3  Related follow-up','Fresh conversation, same arm’s saved code, models, data and notes. Measure correctness, cost, time and repairs again.')],
      takeaway='Controlled access first. Open-web replication is separate and requires an audit of downloaded tools.',
      note='Sandbox probes passed in all three environments: approved HTTPS worked; direct egress, unapproved libraries/data and write requests were blocked. Public ECDS access is not proof of authenticated S2S retrieval.'),
 dict(tag='THREE CASES, ONE EVALUATION PATTERN',title='Numerical evidence underneath. A readable scorecard on top.',
      subtitle='No weighted leaderboard. A scientifically valid negative result can complete the task.',
      rows=[['Case','Independent numerical checks','Scientific measurement'],
            ['Seasonal rainfall','Aggregation, training separation, coverage, probabilities','RPS skill + rainfall RMSE + reliability'],
            ['Forecast IOD','Climatology, weighted index, ensemble spread, dates','Error against OISST; too few windows for calibration'],
            ['Kenya rainfall outlook','Six weekly maps and regional member statistics','Calculation correctness; no observed forecast verification']],
      takeaway='Judge: data/provenance · scientific processing · results/uncertainty · communication · reproduction/reuse. Each rated 0–2 with evidence.',
      note='The judge cannot override failed numerical checks. The main report shows completion, scientific quality, tokens/cost/time, and follow-up performance. Borderline judgments receive human review.'),
 dict(tag='EXISTING FEASIBILITY EVIDENCE — NOT AN ARM COMPARISON',title='A real seasonal problem, with an honest baseline',
      subtitle='September-issued OND rainfall · Kenyan study box · 1993–2008 development years · leave-one-year-out',
      image='../results/development.png',
      cards=[('What the reference calculation shows',f"RPS: raw {SCORES['mean_RPS']['raw']:.3f}; model-climate {SCORES['mean_RPS']['model_climate']:.3f}; calibrated {SCORES['mean_RPS']['calibrated']:.3f}.\n\n15.5% better than raw, but only 1.5% better than model-climate probabilities. Calibration does not uniformly improve."),
             ('What it does not show','These are expert-written feasibility calculations, not agent results. Independent Python reproduced the numerics without ACCORD. We have not measured comparative token cost, speed or reuse.')],
      takeaway='All three arms can in principle solve this task. The experiment tests how reliably and efficiently they do so.',
      note='Exploratory uncertainty intervals include no improvement. Private evaluation years have not been inspected. Source: results/development-scores.json and results/other-arms-feasibility.json.'),
 dict(tag='TODAY’S DELIVERABLE AND REVIEW CHECKPOINT',title='One illustrative suite: three arms, two attempts each',
      subtitle='Prefer the seasonal case to test scientific skill. Kenya is a simpler workflow-efficiency fallback.',
      rows=[['Arm','Initial: correctness / cost / time','Follow-up: correctness / cost / time'],
            ['Scratch','Not run','Not run'],['Rhiza','Not run','Not run'],['ACCORD','Not run','Not run']],
      takeaway='Ready for review: isolated environments, three task contracts, numerical checkers, shared judge rubric and this slide deck.',
      note='PAUSED before model calls. Next: choose model and budget; freeze the chosen reference/data policy; resolve case-specific scientific issues; agree to launch one six-attempt suite. The full three-case pilot is 18 attempts per access mode. No fabricated result ordering.'),
]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG='F5F7F5'; INK='17313B'; TEAL='157C77'; GRAY='58666B'; PALE='E2ECE8'
def text(s,x,y,w,h,value,size=20,color=INK,bold=False):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
 tf.word_wrap=True; tf.margin_left=tf.margin_right=0
 for i,line in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
  p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
 return box

def card(s,x,y,w,h,title,body):
 shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(PALE); shape.line.fill.background()
 text(s,x+.18,y+.17,w-.36,.6,title,20,TEAL,True)
 text(s,x+.18,y+.86,w-.36,h-1,body,18)

for n,data in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
 text(s,.5,.28,12.3,.35,data['tag'],11,TEAL,True)
 text(s,.5,.82,12.3,.7,data['title'],31,INK,True)
 text(s,.5,1.63,12.2,.7,data['subtitle'],18,GRAY)
 if 'rows' in data:
  table=s.shapes.add_table(len(data['rows']),3,Inches(.5),Inches(2.55),Inches(12.3),Inches(2.6)).table
  table.columns[0].width=Inches(2.35); table.columns[1].width=Inches(4.9); table.columns[2].width=Inches(5.05)
  for i,row in enumerate(data['rows']):
   for j,value in enumerate(row):
    cell=table.cell(i,j); cell.text=value; cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor.from_string(TEAL if i==0 else PALE)
    cell.margin_left=Inches(.12); cell.margin_top=Inches(.1)
    for p in cell.text_frame.paragraphs:
     p.font.name='Aptos'; p.font.size=Pt(16); p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK); p.font.bold=i==0
 elif 'image' in data:
  s.shapes.add_picture(str(ROOT/'results/development.png'),Inches(.45),Inches(2.45),width=Inches(7.9))
  title,body=data['cards'][0]; text(s,8.6,2.5,4.2,.5,title,20,TEAL,True); text(s,8.6,3.08,4.15,2.55,body,18)
 else:
  for i,(title,body) in enumerate(data['cards']): card(s,.5+4.17*i,2.55,3.97,2.9,title,body)
 text(s,.5,5.65,12.25,.8,data['takeaway'],20,TEAL,True)
 footer=data['note'] if 'image' not in data else data['cards'][1][1]
 text(s,.5,6.62,11.9,.6,footer,11,GRAY)
 text(s,12.5,7.13,.3,.2,str(n),10,GRAY)
 s.notes_slide.notes_text_frame.text=data['note']
prs.save(ROOT/'slides/accord-weather-pilot.pptx')

sections=[]
for n,s in enumerate(slides,1):
 body=''
 if 'rows' in s:
  body='<table>'+''.join('<tr>'+''.join(('<th>' if i==0 else '<td>')+html.escape(x)+('</th>' if i==0 else '</td>') for x in row)+'</tr>' for i,row in enumerate(s['rows']))+'</table>'
 elif 'image' in s:
  body='<div class="example"><img src="'+s['image']+'"><div>'+''.join('<h3>'+html.escape(t)+'</h3><p>'+html.escape(b).replace('\n','<br>')+'</p>' for t,b in s['cards'])+'</div></div>'
 else:
  body='<div class="cards">'+''.join('<article><h3>'+html.escape(t)+'</h3><p>'+html.escape(b)+'</p></article>' for t,b in s['cards'])+'</div>'
 sections.append('<section><small>'+s['tag']+'</small><h1>'+s['title']+'</h1><p class="sub">'+s['subtitle']+'</p>'+body+'<p class="takeaway">'+s['takeaway']+'</p><footer>'+s['note']+'</footer><span class="num">'+str(n)+'/5</span></section>')
html_page='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ACCORD weather pilot — review deck</title><style>
*{box-sizing:border-box}body{margin:0;background:#d7e2dc;color:#17313b;font:18px system-ui}section{background:#f5f7f5;max-width:1400px;min-height:780px;margin:28px auto;padding:42px 52px;position:relative;box-shadow:0 5px 25px #183a3520}small{color:#157c77;font-weight:700;letter-spacing:1px}h1{font-size:42px;line-height:1.15;margin:26px 0 18px}h3{color:#157c77;font-size:23px}.sub{color:#58666b;font-size:23px}.cards{display:flex;gap:22px;margin:40px 0}.cards article{flex:1;padding:20px;background:#e2ece8;border-radius:12px;min-height:240px}p{line-height:1.55}.takeaway{color:#157c77;font-weight:650;font-size:23px;margin-top:35px}table{width:100%;border-collapse:collapse;margin:40px 0}th,td{padding:20px;text-align:left;border-bottom:2px solid #f5f7f5;background:#e2ece8}th{background:#157c77;color:white}footer{font-size:14px;line-height:1.5;color:#58666b}.num{position:absolute;bottom:20px;right:25px;color:#58666b}.example{display:flex;gap:25px;align-items:center}.example img{width:63%;object-fit:contain}.example p{font-size:16px}.example h3{font-size:20px}@media(max-width:800px){section{padding:24px}h1{font-size:32px}.cards,.example{display:block}.cards article{margin-bottom:15px}.example img{width:100%}table{font-size:13px}th,td{padding:10px}}@media print{@page{size:landscape;margin:0}body{background:white}section{width:100vw;height:100vh;min-height:0;margin:0;padding:28px 40px;box-shadow:none;break-after:page;font-size:15px}h1{font-size:32px}.sub{font-size:18px}.cards{margin:24px 0}.cards article{min-height:200px}.takeaway{font-size:18px}footer{font-size:11px}table{margin:20px 0}th,td{padding:14px}}
</style>'''+''.join(sections)+'''<script>document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown','ArrowLeft','PageUp'].includes(e.key)){let s=[...document.querySelectorAll('section')],i=s.findIndex(x=>x.getBoundingClientRect().bottom>innerHeight/2),d=['ArrowRight','PageDown'].includes(e.key)?1:-1;s[Math.max(0,Math.min(s.length-1,i+d))].scrollIntoView({behavior:'smooth'});e.preventDefault()}})</script></html>'''
(ROOT/'slides/index.html').write_text(html_page)
print('Wrote slides/accord-weather-pilot.pptx and slides/index.html. No agent or judge calls.')
